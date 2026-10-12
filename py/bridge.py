"""Bridge between the tablet web app (JavaScript worker) and the measuring
engine. Every function takes/returns plain JSON-able data.

Same engine as the PC station (core/measure.py, core/training.py)."""
import json
import math

import cv2
import numpy as np

from core.measure import Engine, MeasureError, evaluate
from core import training as T

GEOM = {"thread_major": 10.0, "d12": 12.0, "d14": 14.0, "bend_radius": 14.0,
        "leg2_fit_mm": 25.0}
_engine = None


def engine(geom=None):
    global _engine
    if _engine is None or geom:
        _engine = Engine(dict(GEOM, **(geom or {})))
    return _engine


def _clean(o):
    """numpy -> python, NaN -> None (JSON safe)."""
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_clean(v) for v in o]
    if isinstance(o, np.ndarray):
        return _clean(o.tolist())
    if isinstance(o, (np.floating, float)):
        f = float(o)
        return f if math.isfinite(f) else None
    if isinstance(o, np.integer):
        return int(o)
    return o


def to_gray(buf, w, h, ch):
    if hasattr(buf, "to_py"):            # JS Uint8Array -> memoryview
        buf = buf.to_py()
    a = np.frombuffer(buf, dtype=np.uint8)
    if ch == 1:
        return a.reshape(h, w)
    a = a.reshape(h, w, ch)
    # RGBA from canvas -> grey (same weights as OpenCV)
    return cv2.cvtColor(np.ascontiguousarray(a[:, :, :3]), cv2.COLOR_RGB2GRAY)


def measure(buf, w, h, ch, args_json):
    """args: roi [x,y,w,h] | None, k, model {H,cx,cy} | None, params [...],
    threshold_mode, fixed_threshold"""
    a = json.loads(args_json)
    g = to_gray(buf, w, h, ch)
    out = {"ok": False, "error": None}
    try:
        r = engine().measure(g, a["k"], a.get("roi"), a.get("threshold_mode", "otsu"),
                             a.get("fixed_threshold", 110), True, model=a.get("model"))
        ok, rows = evaluate(r["raw"], a["params"])
        geo = r["geom"]
        cont = geo["contour"]
        step = max(1, len(cont) // 600)
        out.update(ok=ok, rows=rows, raw=r["raw"], features=r["features"],
                   time_ms=r["time_ms"],
                   geom={"contour": cont[::step], "tip": geo["tip"], "far": geo["far"],
                         "P": geo["P"], "corner": geo["corner"],
                         "bend_start_pt": geo["bend_start_pt"],
                         "thread_pts": geo["thread_pts"], "shoulder_pts": geo["shoulder_pts"]})
    except MeasureError as e:
        out["error"] = str(e)
        out["rows"] = evaluate({}, a["params"])[1]
    except Exception as e:                       # never crash the app
        out["error"] = f"Software error: {e}"
        out["rows"] = evaluate({}, a["params"])[1]
    return json.dumps(_clean(out))


def stats_image(buf, w, h, ch):
    """Brightness / saturation / focus of an image (setup check)."""
    g = to_gray(buf, w, h, ch)
    sub = g[::4, ::4]
    cy, cx = g.shape[0] // 2, g.shape[1] // 2
    crop = g[max(0, cy - 250):cy + 250, max(0, cx - 250):cx + 250]
    return json.dumps({"bg": float(np.percentile(sub, 90)),
                       "sat": float((sub >= 254).mean() * 100),
                       "focus": float(cv2.Laplacian(crop, cv2.CV_32F).var())})


def checkerboard(buf, w, h, ch, cols, rows, square_mm):
    g = to_gray(buf, w, h, ch)
    pattern = (int(cols), int(rows))
    ok, c = cv2.findChessboardCornersSB(g, pattern, flags=cv2.CALIB_CB_NORMALIZE_IMAGE) \
        if hasattr(cv2, "findChessboardCornersSB") else (False, None)
    if not ok:
        ok, c = cv2.findChessboardCorners(g, pattern, cv2.CALIB_CB_ADAPTIVE_THRESH)
        if ok:
            cv2.cornerSubPix(g, c, (11, 11), (-1, -1),
                             (cv2.TERM_CRITERIA_EPS + cv2.TERM_CRITERIA_MAX_ITER, 50, 1e-3))
    if not ok:
        return json.dumps({"ok": False})
    p = c.reshape(rows, cols, 2).astype(np.float64)
    dx = np.linalg.norm(np.diff(p, axis=1), axis=2)
    dy = np.linalg.norm(np.diff(p, axis=0), axis=2)
    sx, sy = square_mm / dx.mean(), square_mm / dy.mean()
    k = square_mm / np.concatenate([dx.ravel(), dy.ravel()]).mean()
    return json.dumps({"ok": True, "k": float(k), "tilt_pct": float(abs(sx - sy) / k * 100)})


def scale_from_master(features_json, true_json, k):
    """Quick scale: compare one master measurement with its true values."""
    feats = json.loads(features_json)
    true = json.loads(true_json)
    r = engine().compute(feats, {"k": k})
    rat = [true[x] / r[x] for x in ("d14", "d12", "shoulder", "drop", "thread_len")
           if x in true and r.get(x)]
    if not rat:
        return json.dumps({"ok": False})
    s = float(np.median(rat))
    return json.dumps({"ok": True, "k": k * s, "change_pct": (s - 1) * 100})


def train(caps_json, true_json, params_json, k0, img_w, img_h, model_json, progress=None):
    caps = [{"features": f} for f in json.loads(caps_json)]
    cur = json.loads(model_json) if model_json else None
    res = T.analyse(engine(), caps, json.loads(true_json), json.loads(params_json),
                    float(k0), (int(img_w), int(img_h)), current_model=cur,
                    progress=(lambda t: progress(t)) if progress else None)
    return json.dumps(_clean(res))


def excel(rows_json, params_json, title):
    """Build an .xlsx from the result history. Returns bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import io
    rows = json.loads(rows_json)
    params = [p for p in json.loads(params_json) if p.get("enabled", True)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Inspection"
    ws.append([title])
    ws.append(["No", "Date", "Time", "Operator", "Result", "Failed"] + [p["name"] for p in params])
    for c in ws[2]:
        c.font = Font(bold=True)
    ws.append(["", "", "", "", "", "Limits"] + [
        f"{p['nominal'] - abs(p['tol_minus']):g} – {p['nominal'] + abs(p['tol_plus']):g}"
        for p in params])
    red, green = PatternFill("solid", fgColor="FFC7CE"), PatternFill("solid", fgColor="C6EFCE")
    for i, r in enumerate(rows, 1):
        vals = r.get("values", {})
        ws.append([i, r.get("date"), r.get("time"), r.get("operator", ""),
                   "OK" if r.get("ok") else "NG", r.get("failed", "")] +
                  [None if vals.get(p["key"]) is None else round(vals[p["key"]], 3)
                   for p in params])
        ws.cell(ws.max_row, 5).fill = green if r.get("ok") else red
        for j, p in enumerate(params):
            v = vals.get(p["key"])
            if v is None:
                continue
            lo, hi = p["nominal"] - abs(p["tol_minus"]), p["nominal"] + abs(p["tol_plus"])
            if not (lo - 1e-9 <= v <= hi + 1e-9):
                ws.cell(ws.max_row, 7 + j).fill = red
    n = len(rows)
    ok = sum(1 for r in rows if r.get("ok"))
    s = wb.create_sheet("Summary")
    for line in (["Total", n], ["OK", ok], ["NG", n - ok],
                 ["Yield %", round(ok / n * 100, 2) if n else 0]):
        s.append(line)
    s.append([])
    s.append(["Parameter", "n", "Mean", "Std dev", "Min", "Max", "Cpk"])
    for p in params:
        v = np.array([r["values"][p["key"]] for r in rows
                      if r.get("values", {}).get(p["key"]) is not None], float)
        if len(v) < 2:
            continue
        lo, hi = p["nominal"] - abs(p["tol_minus"]), p["nominal"] + abs(p["tol_plus"])
        sd = v.std(ddof=1)
        cpk = min(hi - v.mean(), v.mean() - lo) / (3 * sd) if sd > 0 else None
        s.append([p["name"], len(v), round(v.mean(), 4), round(sd, 4), round(v.min(), 4),
                  round(v.max(), 4), round(cpk, 2) if cpk is not None else ""])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def demo_image(seed, master, cam_h, k0, w, h):
    """Synthetic camera image (Demo mode without camera). Returns grey bytes."""
    import random
    from core.simpart import render3d, random_part
    rnd = random.Random(int(seed))
    if master:
        from core.simpart import NOMINAL
        g = dict(NOMINAL, thread_len=17.62, shoulder=49.27, bend_start=59.12, d12=11.985,
                 d14=14.03, angle=41.6, drop=37.68)
        ang = rnd.choice([0, 180]) + rnd.uniform(-20, 20)
        off = (rnd.uniform(-1, 1) * max(0, w * k0 / 2 - 62),
               rnd.uniform(-1, 1) * max(0, h * k0 / 2 - 46))
    else:
        g, _ = random_part(0.3, rnd)
        ang = rnd.uniform(-12, 12)
        off = (rnd.uniform(-3, 3), rnd.uniform(-3, 3))
    img, truth = render3d(g, k0, (w, h), ang, off, cam_h=cam_h or None, noise=2.0, ss=2)
    return img.tobytes(), json.dumps(_clean(truth))
